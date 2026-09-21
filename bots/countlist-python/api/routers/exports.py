import csv
import io
import os
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy import select, and_
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from api.config import settings
from api.database.connection import get_db
from api.database.models import Expense, Export, ExportFormat, ExportStatus, User, ExpenseStatus
from api.middleware.auth import get_current_user

router = APIRouter(prefix="/exports", tags=["exports"])


async def _get_expenses(db: AsyncSession, user_id: int, month: Optional[int], year: Optional[int]):
    conditions = [Expense.user_id == user_id, Expense.status == ExpenseStatus.active]
    if month:
        conditions.append(Expense.month == month)
    if year:
        conditions.append(Expense.year == year)
    result = await db.execute(
        select(Expense)
        .options(selectinload(Expense.category))
        .where(and_(*conditions))
        .order_by(Expense.expense_date.desc())
    )
    return result.scalars().all()


@router.get("/csv")
async def export_csv(
    month: Optional[int] = None,
    year: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    expenses = await _get_expenses(db, current_user.id, month, year)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["ID", "Sana", "Miqdor", "Valyuta", "Tavsif", "Kategoriya", "Guruh"])
    for e in expenses:
        writer.writerow([
            e.id,
            e.expense_date.isoformat(),
            e.amount,
            e.currency,
            e.description,
            e.category.name if e.category else "",
            e.group_id or "",
        ])

    output.seek(0)
    filename = f"expenses_{month or 'all'}_{year or 'all'}.csv"
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8-sig")),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/excel")
async def export_excel(
    month: Optional[int] = None,
    year: Optional[int] = None,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    expenses = await _get_expenses(db, current_user.id, month, year)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Xarajatlar"

    headers = ["ID", "Sana", "Miqdor", "Valyuta", "Tavsif", "Kategoriya"]
    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row, e in enumerate(expenses, 2):
        ws.cell(row=row, column=1, value=e.id)
        ws.cell(row=row, column=2, value=e.expense_date.isoformat())
        ws.cell(row=row, column=3, value=float(e.amount))
        ws.cell(row=row, column=4, value=e.currency)
        ws.cell(row=row, column=5, value=e.description)
        ws.cell(row=row, column=6, value=e.category.name if e.category else "")

    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"expenses_{month or 'all'}_{year or 'all'}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
